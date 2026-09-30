import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

def export_exam_to_docx(json_path, output_path):
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    test_name = data.get('TestName', 'វិញ្ញាសាប្រឡង')
    session_name = data.get('SessionName', 'គ្រប់វេនប្រឡងទាំងអស់')
    duration = data.get('DurationMinutes', 45)
    total_marks = data.get('TotalMarks', 100)
    questions = data.get('questions', [])

    doc = docx.Document()
    
    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(test_name)
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0, 40, 142)
    title_p.paragraph_format.space_after = Pt(12)

    # Metadata Table
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    khmer_digits = {'0': '០', '1': '១', '2': '២', '3': '៣', '4': '៤', '5': '៥', '6': '៦', '7': '៧', '8': '៨', '9': '៩'}
    def to_kh(n):
        return ''.join(khmer_digits.get(c, c) for c in str(n))

    row0 = table.rows[0].cells
    row0[0].text = f"វគ្គប្រឡង៖ {session_name}"
    row0[1].text = f"រយៈពេល៖ {to_kh(duration)} នាទី"

    row1 = table.rows[1].cells
    row1[0].text = f"ពិន្ទុសរុប៖ {to_kh(total_marks)} ពិន្ទុ"
    row1[1].text = f"ចំនួនសំណួរ៖ {to_kh(len(questions))} សំណួរ"

    doc.add_paragraph('')

    khmer_letters = ['ក', 'ខ', 'គ', 'ឃ', 'ង', 'ច']

    last_passage = None
    for i, q in enumerate(questions):
        passage = q.get('Passage', q.get('passage', ''))
        if passage and passage.strip() and passage.strip() != (last_passage or '').strip():
            pass_p = doc.add_paragraph()
            pass_p.paragraph_format.space_before = Pt(14)
            pass_p.paragraph_format.space_after = Pt(6)
            p_label = pass_p.add_run("📖 [អត្ថបទអាន]\n")
            p_label.bold = True
            p_label.font.size = Pt(11)
            p_label.font.color.rgb = RGBColor(37, 99, 235)
            
            p_text = pass_p.add_run(passage.strip())
            p_text.font.size = Pt(10.5)
            p_text.italic = True
            p_text.font.color.rgb = RGBColor(51, 65, 85)
            last_passage = passage

        q_text = q.get('QuestionText', q.get('text', ''))
        pts = q.get('Points', q.get('points', 4))
        
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(8)
        qp.paragraph_format.space_after = Pt(2)
        
        q_run = qp.add_run(f"{to_kh(i+1)}. ({to_kh(pts)}ពិន្ទុ) {q_text}")
        q_run.bold = True
        q_run.font.size = Pt(11)

        answers = q.get('answers', [])
        for a_idx, ans in enumerate(answers):
            ans_text = ans.get('AnswerText', ans.get('text', ''))
            is_correct = ans.get('IsCorrect', ans.get('correct', False))
            letter = khmer_letters[a_idx] if a_idx < len(khmer_letters) else chr(65 + a_idx)

            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Inches(0.3)
            ap.paragraph_format.space_before = Pt(0)
            ap.paragraph_format.space_after = Pt(2)

            ans_run = ap.add_run(f"{letter}. {ans_text}")
            ans_run.font.size = Pt(10)
            
            if is_correct:
                ans_run.bold = True
                star_run = ap.add_run("  (ត្រឹមត្រូវ *)")
                star_run.bold = True
                star_run.font.color.rgb = RGBColor(22, 163, 74)

    doc.save(output_path)
    print("SUCCESS")

if __name__ == '__main__':
    if len(sys.argv) >= 3:
        export_exam_to_docx(sys.argv[1], sys.argv[2])
